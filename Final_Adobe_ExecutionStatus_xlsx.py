import requests
import json
import xlsxwriter
import openpyxl

workbook = xlsxwriter.Workbook('chart_bar.xlsx')
worksheet = workbook.add_worksheet()
bold = workbook.add_format({'bold': 1})

#####   Data start #############

Date='2022-12-05';
BaseURL='https://adobe.autonomiq.ai/platform'
username='adobeadmin'
password='Adobeadmin@123'
value=[266,134,73,102,75,139,77,135,299,300,301,298,304,351,372,373,374,375,441,443,444]
Iteration=2;

########### Data Ends ###########

headings = ['TESTCASE NAME','SUCCESS','FAILURE','Execution_ID']

url = "{}/v1/auth".format(BaseURL)

payload = json.dumps({
  "username": username,
  "password": password
})
headers = {
  'Authorization': 'Basic YWRvYmVhZG1pbjpBZG9iZWFkbWluQDEyMw==',
  'Content-Type': 'application/json'
}

response = requests.request("POST", url, headers=headers, data=payload)

j = json.loads(response.text)
token = j['token']

#############  Adding Lists ##############
def sum_lists(*args):
  return list(map(sum, zip(*args)))



#######
count=len(value);
if(Iteration==1):
    TestExecutionName = []
    Failurecount = []
    Successcount = []
    ExecutionId_list = []

    for i in value:
        url = "{}/testScriptExecutions/2/4/{}/-1/0/executions".format(BaseURL,i)

        payload={}
        headers = {
                'Content-Type': 'application/json',
                'content-type': 'multipart/form-data',
                'Authorization': 'Bearer {}'.format(token)
        }

        response = requests.request("GET", url, headers=headers, data=payload)
        L = json.loads(response.text)
        ExecutionCount=L['tasks']

        TestExecutionName.append(ExecutionCount[0]['executionName'])

        if (ExecutionCount[0]['executionStatus'] == "INPROGRESS"):
            Executionvalue = ExecutionCount[1]['executionId']
        else:
            Executionvalue = ExecutionCount[0]['executionId']

        ExecutionId_list.append(Executionvalue)

        failurecount=0;
        successcount=0;
        for i in ExecutionCount:
            date1 = i['initiatedOn']
            date2 = date1.split('T')[0]
            if (Date == date2):
                if (i['executionStatus'] == "ERROR"):
                    failurecount = failurecount + 1;
                if (i['executionStatus'] == "SUCCESS"):
                    successcount = successcount + 1;

        Failurecount.append(failurecount)
        Successcount.append(successcount)

    data = [
        TestExecutionName,
        Successcount,
        Failurecount,
        ExecutionId_list
    ]

    worksheet.write_row('A1', headings, bold)
    worksheet.write_column('A2',data[0])
    worksheet.write_column('B2',data[1])
    worksheet.write_column('C2',data[2])
    worksheet.write_column('D2',data[3])




    chart1 = workbook.add_chart({'type': 'bar'})


    chart1.add_series({
        'name':       ['Sheet1', 0, 2],
        'categories': ['Sheet1', 1, 0, count, 0],
        'values':     ['Sheet1', 1, 2, count, 2],
    })

    chart1.add_series({
        'name':       ['Sheet1', 0, 1],
        'categories': ['Sheet1', 1, 0, count, 0],
        'values':     ['Sheet1', 1, 1, count, 1],
    })

    chart1.set_title({'name': 'Failure data analysis'})

    chart1.set_x_axis({'name': 'Failure count'})

    chart1.set_y_axis({'name': 'Testcase Names(mm)'})

    chart1.set_style(12)

    worksheet.insert_chart('G2', chart1)

    workbook.close()

if(Iteration==2):
    SuccesscountITR_1 = []
    FailurecountITR_1 = []

    SuccesscountITR_2 = []
    FailurecountITR_2 = []
    New_ExecutionId_list =[]
    New_TestExecutionName=[]
    for i in range(0,count):
        print(i)
        j=i+2;
        wb = openpyxl.load_workbook("chart_bar.xlsx")
        sh = wb.active
        s1 = sh['B{}'.format(j)]
        f1 = sh['C{}'.format(j)]
        e1 = sh['D{}'.format(j)]
        SuccesscountITR_1.append(s1.value);
        FailurecountITR_1.append(f1.value)
        url = "{}/testScriptExecutions/2/4/{}/-1/0/executions".format(BaseURL,value[i])

        payload={}
        headers = {
                'Content-Type': 'application/json',
                'content-type': 'multipart/form-data',
                'Authorization': 'Bearer {}'.format(token)
        }

        response = requests.request("GET", url, headers=headers, data=payload)
        L = json.loads(response.text)
        ExecutionCount=L['tasks']

        New_TestExecutionName.append(ExecutionCount[0]['executionName'])

        if (ExecutionCount[0]['executionStatus'] == "INPROGRESS"):
            Executionvalue = ExecutionCount[1]['executionId']
        else:
            Executionvalue = ExecutionCount[0]['executionId']

        New_ExecutionId_list.append(Executionvalue)

        failurecount = 0;
        successcount = 0;
        for i in ExecutionCount:
            if (i['executionId'] == e1.value):
                print(e1.value)
                print(i['executionId'])
                print("inside loop")
                break;
            date1 = i['initiatedOn']
            date2 = date1.split('T')[0]
            if (Date == date2):
                print(i['executionId'])
                if (i['executionStatus'] == "ERROR"):
                    failurecount = failurecount + 1;
                if (i['executionStatus'] == "SUCCESS"):
                    successcount = successcount + 1;

        FailurecountITR_2.append(failurecount)
        SuccesscountITR_2.append(successcount)

    Final_Failure = sum_lists(FailurecountITR_1, FailurecountITR_2)
    Final_Success = sum_lists(SuccesscountITR_1, SuccesscountITR_2)

    print(New_ExecutionId_list)
    print(Final_Failure)
    print(Final_Success)

    data = [
            New_TestExecutionName,
            Final_Success,
            Final_Failure,
            New_ExecutionId_list
        ]
    worksheet.write_row('A1', headings, bold)
    worksheet.write_column('A2', data[0])
    worksheet.write_column('B2', data[1])
    worksheet.write_column('C2', data[2])
    worksheet.write_column('D2', data[3])

    workbook.close()

